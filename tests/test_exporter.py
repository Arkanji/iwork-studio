"""Export — protocol + second-tool verification, with the app stubbed."""

from __future__ import annotations

import hashlib
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))

from iwork_studio import exporter as ex  # noqa: E402
from iwork_studio.keynote_io import read_key  # noqa: E402
from iwork_studio.numbers_io import read_numbers  # noqa: E402


def _sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def _numbers_values(src):
    return [c["value"] for s in read_numbers(src)["sheets"] for t in s["tables"] for c in t["cells"]
            if c.get("value") is not None]


def _key_texts(src):
    return [t for s in read_key(src)["slides"] if s["source"].startswith("Index/Slide") for t in s["texts"]
            if t.strip() and t.strip() != "￼"]


def make(fmt, src, out, password=None, drop=False):
    """What a well-behaved (or broken, drop=True) app would write."""
    if fmt == "pdf":
        from reportlab.lib import pdfencrypt
        from reportlab.pdfgen import canvas

        enc = pdfencrypt.StandardEncryption(password, ownerPassword=password, strength=128) if password else None
        c = canvas.Canvas(str(out), encrypt=enc)
        c.drawString(72, 72, "x")
        c.save()
    elif fmt == "xlsx":
        import openpyxl

        wb = openpyxl.Workbook()
        vals = _numbers_values(src)
        for i, v in enumerate(vals[:-1] if drop else vals, start=1):
            wb.active.cell(i, 1, v)
        wb.save(out)
    elif fmt == "csv":
        vals = _numbers_values(src)
        Path(out).write_text("\n".join(str(v) for v in (vals[:-1] if drop else vals)), encoding="utf-8")
    elif fmt == "pptx":
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        for t in ([] if drop else _key_texts(src)):
            sl.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = t
        prs.save(out)
    elif fmt == "images":
        Path(out).mkdir()
        (Path(out) / "1.jpeg").write_bytes(b"\xff\xd8\xff")


@pytest.fixture()
def key_file(tmp_path):
    d = tmp_path / "deck.key"
    d.write_bytes((REPO / "tests" / "fixtures" / "arabic.key").read_bytes())
    return d


@pytest.fixture()
def fake_app(monkeypatch):
    state = {"drop": False, "touch_source": False, "calls": []}

    def fake(src, out, kind, as_name, props):
        state["calls"].append((kind, as_name, props))
        make(as_name_to_fmt[as_name], src, out, props.get("password"), state["drop"])
        if state["touch_source"]:
            src.write_bytes(src.read_bytes() + b"\0")

    as_name_to_fmt = {v[0]: k for fmts in ex.FORMATS.values() for k, v in fmts.items()}
    monkeypatch.setattr(ex, "_jxa_export", fake)
    return state


def test_xlsx_export_checked_against_numbers_values(numbers_file, fake_app):
    r = ex.export(numbers_file, "xlsx")
    assert r["ok"] and Path(r["output"]).suffix == ".xlsx" and r["source_unchanged"]


def test_csv_with_arabic(numbers_file, fake_app):
    r = ex.export(numbers_file, "csv")
    assert "الاسم" in Path(r["output"]).read_text(encoding="utf-8")


def test_missing_value_rejected_nothing_written(numbers_file, fake_app):
    fake_app["drop"] = True
    with pytest.raises(ex.ExportError, match="missing source values"):
        ex.export(numbers_file, "xlsx")
    assert not numbers_file.with_suffix(".xlsx").exists()


def test_pptx_export_checked_against_keynote_text(key_file, fake_app):
    assert ex.export(key_file, "pptx")["checked"]["slides"] == 1
    fake_app["drop"] = True
    with pytest.raises(ex.ExportError, match="missing slide text"):
        ex.export(key_file, "pptx", overwrite=True)


def test_password_pdf_must_be_encrypted(numbers_file, fake_app):
    r = ex.export(numbers_file, "pdf", password="s3cret", password_hint="usual")
    assert r["checked"]["encrypted"]
    assert fake_app["calls"][-1][2] == {"password": "s3cret", "passwordHint": "usual"}


def test_source_changed_during_export_is_rejected(numbers_file, fake_app):
    fake_app["touch_source"] = True
    with pytest.raises(ex.ExportError, match="changed during export"):
        ex.export(numbers_file, "pdf")
    assert not numbers_file.with_suffix(".pdf").exists()


def test_existing_destination_refused_then_backed_up(numbers_file, fake_app):
    dest = numbers_file.with_suffix(".pdf")
    dest.write_bytes(b"old")
    with pytest.raises(ex.ExportError, match="already exists"):
        ex.export(numbers_file, "pdf")
    r = ex.export(numbers_file, "pdf", overwrite=True)
    assert Path(r["previous_output_saved_as"]).read_bytes() == b"old"


def test_slide_images_folder(key_file, fake_app, monkeypatch):
    converted = []
    monkeypatch.setattr(ex, "_convert_images", lambda folder, fmt: converted.append(fmt))
    r = ex.export(key_file, "images", image_format="png")
    assert Path(r["output"]).is_dir() and r["checked"]["images"] == 1
    # Keynote's JXA export rejects imageFormat (-1700): never sent, converted afterwards
    assert "imageFormat" not in fake_app["calls"][-1][2] and converted == ["png"]


def test_convert_images_skips_matching_format(tmp_path, monkeypatch):
    (tmp_path / "s.001.png").write_bytes(b"png")
    calls = []
    monkeypatch.setattr(ex.subprocess, "run", lambda *a, **k: calls.append(a))
    ex._convert_images(tmp_path, "png")
    assert calls == [] and (tmp_path / "s.001.png").exists()


@pytest.mark.parametrize("call", [
    lambda f: ex.export(f, "pptx"),                         # Numbers can't do pptx
    lambda f: ex.export(f, "csv", password="x"),            # no password for csv
    lambda f: ex.export(f, "pdf", image_quality="ultra"),
    lambda f: ex.export(f, "pdf", out=str(f)),              # onto the source
])
def test_bad_requests(numbers_file, fake_app, call):
    before = _sha(numbers_file)
    with pytest.raises(ex.ExportError):
        call(numbers_file)
    assert _sha(numbers_file) == before and fake_app["calls"] == []


@pytest.mark.aqua
@pytest.mark.parametrize("fmt", ["pdf", "xlsx", "csv"])
def test_live_numbers_export(numbers_file, fmt):
    assert ex.export(numbers_file, fmt)["ok"]


@pytest.mark.aqua
@pytest.mark.parametrize("fmt", ["pdf", "pptx", "images"])
def test_live_keynote_export(key_file, fmt):
    assert ex.export(key_file, fmt)["ok"]


def test_cellnorm_uses_spreadsheet_precision():
    # an xlsx keeps 33.999999999999996 (an older numbers-parser write of 34); the source reads 34
    assert ex._cellnorm(33.999999999999996) == ex._cellnorm(34.0) == "34"
    assert ex._cellnorm(0.1 + 0.2) == ex._cellnorm(0.3)
    assert ex._cellnorm(3.14) == "3.14"


@pytest.mark.aqua
def test_live_slide_images_as_jpeg(tmp_path):
    src = tmp_path / "d.key"
    src.write_bytes((REPO / "tests" / "fixtures" / "arabic.key").read_bytes())
    r = ex.export(src, "images", tmp_path / "slides", image_format="jpeg")
    files = [f for f in Path(r["output"]).rglob("*") if f.is_file()]
    assert files and all(f.suffix == ".jpeg" for f in files) and files[0].read_bytes()[:2] == b"\xff\xd8"
